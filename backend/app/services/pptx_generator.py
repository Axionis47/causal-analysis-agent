"""PowerPoint report generation service."""

from __future__ import annotations

import base64
import io
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from app.core.config import settings
from app.core.logging import get_logger
from app.services.storage import ensure_local_dir, local_storage_root, upload_file

logger = get_logger(__name__)


def _format_number(value: Any, decimals: int = 4) -> str:
    """Format a number for display."""
    if value is None:
        return "N/A"
    try:
        return f"{float(value):.{decimals}f}"
    except (TypeError, ValueError):
        return str(value)


def generate_pptx_report(
    analysis_id: uuid.UUID,
    dataset_name: str | None,
    executive_summary: dict[str, Any],
    technical_report: dict[str, Any],
    effects: list[dict[str, Any]],
    graphs: list[dict[str, Any]],
    validation_results: list[dict[str, Any]],
    visualizations: dict[str, Any] | None = None,
) -> Path | None:
    """
    Generate a PowerPoint report for an analysis.

    Args:
        analysis_id: UUID of the analysis
        dataset_name: Name of the dataset (e.g., Kaggle URL)
        executive_summary: Executive summary data
        technical_report: Technical report data
        effects: List of treatment effects
        graphs: List of causal graphs
        validation_results: List of validation test results
        visualizations: Optional visualization data including chart images

    Returns:
        Path to the generated PPTX file, or None if generation failed
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RgbColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        logger.warning("python-pptx not installed, skipping PPTX generation")
        return None

    # Create reports directory
    root = local_storage_root() / str(analysis_id) / "reports"
    ensure_local_dir(root)
    pptx_path = root / "report.pptx"

    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Define PlotPointe branding colors
    accent_color = RgbColor(230, 103, 59)  # #E6673B

    # --- Slide 1: Title ---
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(12), Inches(1.5))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "Causal Analysis Report"
    title_para.font.size = Pt(44)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER

    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12), Inches(0.5))
    subtitle_frame = subtitle_box.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.text = f"Analysis ID: {str(analysis_id)[:8]}..."
    subtitle_para.font.size = Pt(18)
    subtitle_para.font.color.rgb = RgbColor(128, 128, 128)
    subtitle_para.alignment = PP_ALIGN.CENTER

    if dataset_name:
        dataset_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12), Inches(0.5))
        dataset_frame = dataset_box.text_frame
        dataset_para = dataset_frame.paragraphs[0]
        dataset_para.text = f"Dataset: {dataset_name}"
        dataset_para.font.size = Pt(14)
        dataset_para.font.color.rgb = RgbColor(128, 128, 128)
        dataset_para.alignment = PP_ALIGN.CENTER

    date_box = slide.shapes.add_textbox(Inches(0.5), Inches(5), Inches(12), Inches(0.5))
    date_frame = date_box.text_frame
    date_para = date_frame.paragraphs[0]
    date_para.text = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    date_para.font.size = Pt(14)
    date_para.font.color.rgb = RgbColor(128, 128, 128)
    date_para.alignment = PP_ALIGN.CENTER

    # --- Slide 2: Executive Summary ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _add_slide_title(slide, "Executive Summary")

    overview = executive_summary.get("overview", "Analysis completed successfully.")
    overview_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(1.5))
    overview_frame = overview_box.text_frame
    overview_frame.word_wrap = True
    overview_para = overview_frame.paragraphs[0]
    overview_para.text = overview[:500] if len(overview) > 500 else overview
    overview_para.font.size = Pt(14)

    # Key findings
    findings = executive_summary.get("key_findings", [])
    if findings:
        findings_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12), Inches(3))
        findings_frame = findings_box.text_frame
        findings_frame.word_wrap = True

        findings_title = findings_frame.paragraphs[0]
        findings_title.text = "Key Findings"
        findings_title.font.bold = True
        findings_title.font.size = Pt(16)

        for finding in findings[:5]:  # Limit to 5 findings
            p = findings_frame.add_paragraph()
            p.text = f"• {finding[:100]}"
            p.font.size = Pt(12)
            p.level = 0

    # --- Slide 3: Treatment Effects ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _add_slide_title(slide, "Treatment Effects")

    if effects:
        # Create a simple text table
        table_top = Inches(1.5)
        table_left = Inches(0.5)
        col_widths = [Inches(3), Inches(3), Inches(2), Inches(2), Inches(2)]
        row_height = Inches(0.4)

        # Header row
        headers = ["Treatment", "Outcome", "Method", "ATE", "CI"]
        for col_idx, (header, width) in enumerate(zip(headers, col_widths)):
            x = table_left + sum(col_widths[:col_idx])
            header_box = slide.shapes.add_textbox(x, table_top, width, row_height)
            header_frame = header_box.text_frame
            header_para = header_frame.paragraphs[0]
            header_para.text = header
            header_para.font.bold = True
            header_para.font.size = Pt(11)

        # Data rows
        for row_idx, effect in enumerate(effects[:8]):  # Limit to 8 rows
            y = table_top + (row_idx + 1) * row_height
            values = [
                effect.get("treatment", "N/A")[:20],
                effect.get("outcome", "N/A")[:20],
                effect.get("method", "N/A")[:15],
                _format_number(effect.get("ate"), 4),
                f"[{_format_number(effect.get('ci_lower', effect.get('ate_ci_lower')), 2)}, "
                f"{_format_number(effect.get('ci_upper', effect.get('ate_ci_upper')), 2)}]",
            ]
            for col_idx, (value, width) in enumerate(zip(values, col_widths)):
                x = table_left + sum(col_widths[:col_idx])
                cell_box = slide.shapes.add_textbox(x, y, width, row_height)
                cell_frame = cell_box.text_frame
                cell_para = cell_frame.paragraphs[0]
                cell_para.text = value
                cell_para.font.size = Pt(10)
    else:
        no_data = slide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(12), Inches(1))
        no_data.text_frame.paragraphs[0].text = "No treatment effects available."

    # --- Slide 4: Validation Results ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _add_slide_title(slide, "Validation Results")

    if validation_results:
        y_pos = Inches(1.5)
        for result in validation_results[:6]:  # Limit to 6 results
            method = result.get("method", "Unknown")
            passed = result.get("passed", False)
            confidence = result.get("confidence_score")

            status = "PASSED" if passed else "FAILED"
            status_color = RgbColor(34, 197, 94) if passed else RgbColor(239, 68, 68)

            result_box = slide.shapes.add_textbox(Inches(0.5), y_pos, Inches(10), Inches(0.5))
            result_frame = result_box.text_frame
            p = result_frame.paragraphs[0]
            p.text = f"{method}: "
            p.font.size = Pt(14)

            run = p.add_run()
            run.text = status
            run.font.bold = True
            run.font.color.rgb = status_color

            if confidence is not None:
                run = p.add_run()
                run.text = f" (Confidence: {confidence * 100:.0f}%)"
                run.font.size = Pt(14)

            y_pos += Inches(0.6)
    else:
        no_data = slide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(12), Inches(1))
        no_data.text_frame.paragraphs[0].text = "No validation results available."

    # --- Slide 5: Causal Graph Summary ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _add_slide_title(slide, "Causal Graph")

    if graphs:
        graph = graphs[0]
        nodes = graph.get("nodes", [])
        edges = graph.get("edges", [])
        method = graph.get("method", "Unknown")
        confidence = graph.get("confidence")

        info_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(6), Inches(2))
        info_frame = info_box.text_frame
        info_frame.word_wrap = True

        p = info_frame.paragraphs[0]
        p.text = f"Discovery Method: {method}"
        p.font.size = Pt(14)

        p = info_frame.add_paragraph()
        p.text = f"Nodes: {len(nodes)}"
        p.font.size = Pt(14)

        p = info_frame.add_paragraph()
        p.text = f"Edges: {len(edges)}"
        p.font.size = Pt(14)

        if confidence:
            p = info_frame.add_paragraph()
            p.text = f"Confidence: {confidence * 100:.0f}%"
            p.font.size = Pt(14)

        # List nodes
        if nodes:
            nodes_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12), Inches(2.5))
            nodes_frame = nodes_box.text_frame
            nodes_frame.word_wrap = True

            p = nodes_frame.paragraphs[0]
            p.text = "Variables in Graph:"
            p.font.bold = True
            p.font.size = Pt(12)

            node_names = [n.get("name", str(n)) if isinstance(n, dict) else str(n) for n in nodes]
            p = nodes_frame.add_paragraph()
            p.text = ", ".join(node_names[:15])
            p.font.size = Pt(11)
            if len(node_names) > 15:
                p = nodes_frame.add_paragraph()
                p.text = f"... and {len(node_names) - 15} more"
                p.font.size = Pt(11)
                p.font.italic = True
    else:
        no_data = slide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(12), Inches(1))
        no_data.text_frame.paragraphs[0].text = "No causal graph available."

    # --- Slide 6: Recommendations ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _add_slide_title(slide, "Recommendations")

    recommendations = executive_summary.get("recommendations", [])
    if recommendations:
        rec_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5))
        rec_frame = rec_box.text_frame
        rec_frame.word_wrap = True

        for rec in recommendations[:6]:
            p = rec_frame.add_paragraph()
            p.text = f"• {rec}"
            p.font.size = Pt(14)
            p.level = 0
    else:
        rec_box = slide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(12), Inches(2))
        rec_frame = rec_box.text_frame

        default_recs = [
            "Review the causal graph with domain experts",
            "Validate treatment effects with additional data",
            "Consider sensitivity analysis for key findings",
        ]
        for rec in default_recs:
            p = rec_frame.add_paragraph()
            p.text = f"• {rec}"
            p.font.size = Pt(14)

    # Save presentation
    try:
        prs.save(str(pptx_path))
        logger.info(
            "PPTX report generated",
            analysis_id=str(analysis_id),
            file_size=pptx_path.stat().st_size,
        )
        return pptx_path
    except Exception as exc:
        logger.exception("Failed to save PPTX report", error=str(exc))
        return None


def _add_slide_title(slide: Any, title: str) -> None:
    """Add a title to a slide."""
    from pptx.util import Inches, Pt

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(28)
    title_para.font.bold = True


async def generate_and_upload_pptx(
    analysis_id: uuid.UUID,
    dataset_name: str | None,
    executive_summary: dict[str, Any],
    technical_report: dict[str, Any],
    effects: list[dict[str, Any]],
    graphs: list[dict[str, Any]],
    validation_results: list[dict[str, Any]],
    bucket_name: str | None = None,
) -> tuple[Path | None, str | None]:
    """
    Generate PPTX report and optionally upload to GCS.

    Returns:
        Tuple of (local_path, gcs_path)
    """
    pptx_path = generate_pptx_report(
        analysis_id=analysis_id,
        dataset_name=dataset_name,
        executive_summary=executive_summary,
        technical_report=technical_report,
        effects=effects,
        graphs=graphs,
        validation_results=validation_results,
    )

    if pptx_path is None:
        return None, None

    gcs_path = None
    if bucket_name:
        try:
            gcs_path = await upload_file(
                pptx_path,
                bucket_name=bucket_name,
                destination=f"{analysis_id}/report.pptx",
            )
        except Exception as exc:
            logger.warning("Failed to upload PPTX to GCS", error=str(exc))

    return pptx_path, gcs_path
